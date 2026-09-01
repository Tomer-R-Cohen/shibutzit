"""Generate a small synthetic student workbook for tests and demos.

Produces the small legacy synthetic_students.xlsx fixture and a larger
conversational_demo_students.xlsx fixture. The latter contains explicit
support categories and friendship requests for realistic conversational
tradeoff demos. Both use clearly fake identities and the same physical shape
as the real workbook (header on row 4, data starting row 5).
"""

from __future__ import annotations

import os

import openpyxl

FAKE_FIRST_NAMES = [
    "דוגמנית", "בדיקנית", "נסיינית", "לדוגמה", "טסטונית", "פקטית", "מוקנית",
    "דמיונית", "רגילה", "מיוחדת", "כללית", "פרטית", "משנית", "ראשית",
    "עזרנית", "חדשה", "ותיקה", "שנייה", "שלישית", "רביעית", "חמישית",
    "שישית", "שביעית", "שמינית", "תשיעית", "עשירית", "אחת עשרה", "שתים עשרה",
]
FAKE_LAST_NAMES = [
    "דוגמה", "בדיקה", "טסט", "נתון", "פקט", "מבחן", "סימולציה", "ניסוי",
]
SCHOOLS = ["בית ספר דוגמה א", "בית ספר דוגמה ב", "בית ספר דוגמה ג"]
LEVELS = ["מצטיינת", "בינונית", "חלשה"]


def build_rows(n: int = 28) -> list[dict]:
    rows = []
    for i in range(1, n + 1):
        rows.append(
            {
                "idx": i,
                "last": FAKE_LAST_NAMES[i % len(FAKE_LAST_NAMES)],
                "first": FAKE_FIRST_NAMES[(i - 1) % len(FAKE_FIRST_NAMES)],
                "school": SCHOOLS[i % len(SCHOOLS)],
                "cur_class": i % 5,
                "origin": "א" if i % 4 == 0 else None,
                "level": LEVELS[i % 3],
            }
        )
    return rows


def write_workbook(path: str) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.cell(row=2, column=4, value="רשימה סינתטית לבדיקות - שכבת ז'")
    headers = ["", "מספר סידורי", "שם משפחה", "שם פרטי", 'ביה"ס נוכחי', "כיתה", "מוצא", "הישגים לימודיים"]
    for col, h in enumerate(headers, start=1):
        ws.cell(row=4, column=col, value=h)

    rows = build_rows(28)
    r = 5
    for i, row in enumerate(rows):
        ws.cell(row=r, column=2, value=row["idx"])
        ws.cell(row=r, column=3, value=row["last"])
        ws.cell(row=r, column=4, value=row["first"])
        ws.cell(row=r, column=5, value=row["school"])
        ws.cell(row=r, column=6, value=row["cur_class"])
        ws.cell(row=r, column=7, value=row["origin"])
        ws.cell(row=r, column=8, value=row["level"])
        r += 1
        if i == 10:
            r += 1  # insert a blank spacer row, matching real workbook quirk

    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb.save(path)


def build_conversational_rows(n: int = 72) -> list[dict]:
    """A feasible roster whose social clusters compete with balance goals."""
    schools = [f"בית ספר מקור {letter}" for letter in "אבגדהו"]
    rows = []
    for i in range(1, n + 1):
        first = "תלמידה"
        last = f"מדומה {i:02d}"
        # Academic levels and source schools occur in blocks. Friendship
        # rings also stay inside six-student blocks, creating a real tradeoff:
        # keeping every ring intact is at odds with spreading levels/schools.
        level = LEVELS[(i - 1) // 24]
        school = schools[(i - 1) // 12]
        group_start = ((i - 1) // 6) * 6 + 1
        position = (i - group_start) % 6
        friend_ids = [
            group_start + ((position - 1) % 6),
            group_start + ((position + 1) % 6),
            ((i + 17 - 1) % n) + 1,
        ]
        friend_names = [f"תלמידה מדומה {friend_id:02d}" for friend_id in friend_ids]
        rows.append(
            {
                "idx": i,
                "last": last,
                "first": first,
                "school": school,
                "cur_class": ((i - 1) % 6) + 1,
                "origin": "א" if 28 <= i <= 45 else None,
                "level": level,
                "differential": "כן" if i <= 6 else None,
                "inclusion": "כן" if 7 <= i <= 18 else None,
                "hamar": "כן" if 19 <= i <= 27 else None,
                "friends": ", ".join(friend_names),
            }
        )
    return rows


def write_conversational_workbook(path: str) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "תלמידות"
    ws.sheet_view.rightToLeft = True
    ws.cell(row=2, column=2, value="נתוני הדגמה אנונימיים לשיחה ואופטימיזציה")
    headers = [
        "מספר סידורי",
        "שם משפחה",
        "שם פרטי",
        'ביה"ס נוכחי',
        "כיתה",
        "מוצא",
        "הישגים לימודיים",
        "תלמידה דיפרנציאלית",
        "תלמידה בשילוב",
        'סטטוס ח"מ',
        "בקשות חברות (שמות, מופרד בפסיקים)",
    ]
    for column, header in enumerate(headers, start=1):
        ws.cell(row=4, column=column, value=header)

    target_row = 5
    for index, row in enumerate(build_conversational_rows(), start=1):
        values = [
            row["idx"], row["last"], row["first"], row["school"], row["cur_class"],
            row["origin"], row["level"], row["differential"], row["inclusion"],
            row["hamar"], row["friends"],
        ]
        for column, value in enumerate(values, start=1):
            ws.cell(row=target_row, column=column, value=value)
        target_row += 1
        if index in (24, 48):
            target_row += 1

    ws.freeze_panes = "A5"
    widths = {1: 14, 2: 18, 3: 14, 4: 22, 5: 10, 6: 10, 7: 20, 8: 22, 9: 18, 10: 14, 11: 70}
    for column, width in widths.items():
        ws.column_dimensions[openpyxl.utils.get_column_letter(column)].width = width
    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb.save(path)


if __name__ == "__main__":
    write_workbook(os.path.join(os.path.dirname(__file__), "synthetic_students.xlsx"))
    write_conversational_workbook(os.path.join(os.path.dirname(__file__), "conversational_demo_students.xlsx"))
    print("wrote synthetic_students.xlsx and conversational_demo_students.xlsx")
